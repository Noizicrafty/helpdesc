from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ProcessedEmail
from .utils import ensure_directory, format_datetime


INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
MAX_SHEET_TITLE_LENGTH = 31


class ExportService:
    def export(
        self,
        processed_emails: List[ProcessedEmail],
        destination: str | Path,
        include_replies: bool,
        include_heatmaps: bool,
    ) -> Path:
        """Экспортирует результат обработки в Excel-книгу.

        В отличие от CSV, формат XLSX позволяет создать несколько листов в одном
        выходном документе. Поэтому каждая тема/категория сохраняется на отдельном
        листе, а параметры писем размещаются по отдельным столбцам.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        export_dir = ensure_directory(Path(destination))
        export_path = export_dir / f"mail_export_{timestamp}.xlsx"

        grouped: Dict[str, List[ProcessedEmail]] = defaultdict(list)
        for item in processed_emails:
            grouped[item.assignment.category].append(item)

        workbook = Workbook()
        default_sheet = workbook.active

        if not grouped:
            default_sheet.title = "Без данных"
            default_sheet.append(["Нет обработанных писем для экспорта"])
            workbook.save(export_path)
            return export_path

        workbook.remove(default_sheet)
        used_titles: set[str] = set()

        for category in sorted(grouped.keys(), key=lambda value: value.lower()):
            sheet_title = self._safe_sheet_title(category, used_titles)
            worksheet = workbook.create_sheet(title=sheet_title)
            self._fill_category_sheet(
                worksheet=worksheet,
                category=category,
                items=sorted(grouped[category], key=lambda item: item.email.received_at),
                include_replies=include_replies,
                include_heatmaps=include_heatmaps,
            )

        workbook.save(export_path)
        return export_path

    def _safe_sheet_title(self, raw_title: str, used_titles: set[str]) -> str:
        title = INVALID_SHEET_CHARS.sub(" ", raw_title or "Тема")
        title = re.sub(r"\s+", " ", title).strip() or "Тема"
        base = title[:MAX_SHEET_TITLE_LENGTH]
        candidate = base
        index = 2

        while candidate in used_titles:
            suffix = f"_{index}"
            candidate = f"{base[:MAX_SHEET_TITLE_LENGTH - len(suffix)]}{suffix}"
            index += 1

        used_titles.add(candidate)
        return candidate

    def _fill_category_sheet(
        self,
        worksheet,
        category: str,
        items: List[ProcessedEmail],
        include_replies: bool,
        include_heatmaps: bool,
    ) -> None:
        headers = [
            "№",
            "Отправитель",
            "Email отправителя",
            "Тема письма",
            "Текст письма",
            "Дата получения",
            "Категория",
            "Уверенность",
            "Похоже на спам/рекламу",
            "Spam score",
            "Причина выбора категории",
        ]
        if include_heatmaps:
            headers.append("Ключевые термы")
        if include_replies:
            headers.append("Предлагаемый ответ")

        worksheet.append(headers)
        self._style_header(worksheet)

        for row_number, item in enumerate(items, start=1):
            email = item.email
            row = [
                row_number,
                email.sender_name,
                email.sender_email,
                email.subject,
                email.normalized_body or email.body,
                format_datetime(email.received_at),
                category,
                item.assignment.confidence,
                "Да" if email.is_spam_like else "Нет",
                email.spam_score,
                item.assignment.category_reason,
            ]

            if include_heatmaps:
                terms = "; ".join(
                    f"{term}: {value}"
                    for term, value in zip(item.assignment.heatmap_terms, item.assignment.heatmap_values)
                )
                row.append(terms)

            if include_replies:
                row.append(item.suggested_reply or "")

            worksheet.append(row)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        self._fit_columns(worksheet)
        self._style_body(worksheet)

    def _style_header(self, worksheet) -> None:
        fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _style_body(self, worksheet) -> None:
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _fit_columns(self, worksheet) -> None:
        max_width_by_column = {
            "A": 6,
            "B": 24,
            "C": 28,
            "D": 34,
            "E": 70,
            "F": 20,
            "G": 24,
            "H": 14,
            "I": 22,
            "J": 12,
            "K": 42,
            "L": 34,
            "M": 70,
        }

        for column_index in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column_index)
            worksheet.column_dimensions[letter].width = max_width_by_column.get(letter, 24)
